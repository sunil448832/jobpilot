#!/usr/bin/env python3
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

SRC = "job-tracker.csv"
OUT = "job-tracker.xlsx"

with open(SRC, newline="", encoding="utf-8") as f:
    rows = list(csv.reader(f))
header, data = rows[0], rows[1:]

wb = Workbook()
ws = wb.active
ws.title = "Jobs"

# Styles
head_fill = PatternFill("solid", fgColor="1F4E78")
head_font = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(vertical="top", wrap_text=True)

# Header
ws.append(header)
for c in ws[1]:
    c.fill = head_fill
    c.font = head_font
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    c.border = border

# Data rows
for r in data:
    ws.append(r)

# Alternating row shading + borders + wrap
alt = PatternFill("solid", fgColor="F2F6FB")
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
    for c in row:
        c.border = border
        c.alignment = wrap
        if i % 2 == 0:
            c.fill = alt

# Column widths
widths = {"Company":18,"Role":40,"Location":16,"Region":10,"Work Mode":12,
          "Source":13,"Job Link":32,"Expected/Posted Salary":24,"Date Applied":13,
          "Stage":16,"Next Action":22,"Follow-up Date":14,"Contact/Recruiter":18,
          "Referral?":10,"Resume Version":14,"Priority":10,"Notes":34}
from openpyxl.utils import get_column_letter
for idx, name in enumerate(header, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 16)

# Freeze header + first two cols
ws.freeze_panes = "C2"

# Dropdown validations
def add_dv(col_name, options):
    if col_name not in header:
        return
    ci = header.index(col_name) + 1
    letter = get_column_letter(ci)
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("%s2:%s500" % (letter, letter))

add_dv("Stage", ["Interested","To Apply","Applied","Recruiter Screen",
                 "Technical/Case","Onsite/Final","Offer","Negotiating",
                 "Accepted","Rejected","Withdrawn","Ghosted"])
add_dv("Priority", ["High","Medium","Low"])
add_dv("Region", ["India","UAE","Remote","Other"])
add_dv("Work Mode", ["On-site","Hybrid","Remote"])
add_dv("Referral?", ["Yes","No"])

ws.auto_filter.ref = ws.dimensions

wb.save(OUT)
print("wrote", OUT, "with", len(data), "job rows")
