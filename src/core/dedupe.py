#!/usr/bin/env python3
"""
dedupe.py — one pass that de-duplicates every store in the pipeline.

Duplicates have shown up in every layer, each time for a different reason, and
each time they were fixed locally and came back somewhere else:

  boards.yaml    a slug listed twice in CANDIDATES; the same Phenom host
                 registered under two locale paths (/inception and /global/en)
  state.db       the same req posted in 7 cities = 7 rows (Databricks FDE), so a
                 screen verdict on one left the other six live
  queue/         re-filling an application wrote a second item, orphaning the
                 drafted answers in the first
  job-tracker    a submitted role added again from the ranked list
  referral       the same person tracked twice for one role

So this runs as a routine step, not as an incident response. It is idempotent and
always keeps the RICHEST record — most jobs, most fields, furthest status — never
merely the newest.

    python jobs/dedupe.py            # clean everything
    python jobs/dedupe.py --dry-run
    python jobs/dedupe.py --only boards|jobs|queue|tracker|referrals
"""
import argparse
import glob
import json
import os
import re
import sqlite3
import sys

from jobpilot.core.paths import (SRC as JOBS_DIR, TOOL, CONFIG, DATA, TRACKING, POLICY,  # noqa: E402
                   RESUME, APPLICATIONS, TRACKERS, MEMORY)
DB = os.path.join(DATA, "state.db")
QUEUE = os.path.join(DATA, "queue")
BOARDS = os.path.join(CONFIG, "boards.yaml")

# how far along a queue item is — never drop a further one for a fresher one
RANK = {"submitted": 6, "unconfirmed": 6, "submitting": 5, "approved": 4, "needs_input": 3, "manual": 3,
        "pending": 2, "failed": 1, "rejected": 0, "superseded": -1}
# a card that was, or may have been, sent is never removed: it is what stops a second filing
SENT = ("submitted", "unconfirmed", "submitting")

JUNK_SLUGS = {"assets", "assets-aws", "static", "cdn", "media", "img", "www",
              "api", "boards", "boards-api", "job-boards", "embed", "widget"}


def norm(s):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]+", " ", (s or "").lower())).strip()


def market_priorities():
    """{market_id: priority} straight from targets.yaml; lower is better."""
    try:
        import yaml
        t = yaml.safe_load(open(os.path.join(CONFIG, "targets.yaml"))) or {}
        pri = {m["id"]: m.get("priority", 9) for m in t.get("markets", [])}
        return pri or {"usa": 9}
    except Exception:
        # Never let a malformed targets.yaml turn dedupe into a no-op; falling
        # back to "everything ties" still keeps the best score and fullest JD.
        return {}


# ---------------------------------------------------------------- boards.yaml

def dedupe_boards(dry=False):
    import yaml
    if not os.path.isfile(BOARDS):
        return 0
    d = yaml.safe_load(open(BOARDS)) or {}
    before = sum(len(v) for v in d.values())
    for plat in list(d):
        best = {}
        for b in d[plat]:
            slug = (b.get("slug") or "").lower()
            if not slug or slug in JUNK_SLUGS:
                continue
            # A Phenom host found under two locale paths is ONE board.
            prev = best.get(slug)
            if prev is None or b.get("jobs", 0) > prev.get("jobs", 0):
                best[slug] = b
        d[plat] = sorted(best.values(), key=lambda b: b["slug"])
        if not d[plat]:
            del d[plat]
    after = sum(len(v) for v in d.values())
    if not dry and after != before:
        yaml.safe_dump(d, open(BOARDS, "w"), sort_keys=False)
    return before - after


# -------------------------------------------------------------------- state.db

def dedupe_jobs(dry=False):
    if not os.path.isfile(DB):
        return 0
    con = sqlite3.connect(DB)
    rows = con.execute("SELECT key, company, title, market, score, status, "
                       "LENGTH(COALESCE(jd,'')) FROM jobs").fetchall()
    # Same company+title in several cities is one role. Keep the copy in the best
    # market, then the best score, then the fullest JD.
    #
    # Read from targets.yaml rather than hardcoded. The hardcoded map listed only
    # the original six markets, so every market added later fell to the default
    # and ranked WORSE than the USA — OpenAI "Applied AI Engineer, Cyber" scored
    # 94.4 as a Dublin role, then dedupe threw that copy away and kept the 76.8
    # US one. Deriving it means adding a market can never desync this again.
    PRIORITY = market_priorities()
    groups = {}
    for key, co, title, market, score, status, jdlen in rows:
        groups.setdefault((norm(co), norm(title)[:60]), []).append(
            # 99, not 4: with priorities read from targets.yaml the USA is now
            # itself 4, so a 4 default would tie with it instead of losing.
            (key, PRIORITY.get(market, 99), -(score or 0), -(jdlen or 0), status))
    drop = []
    for g in groups.values():
        if len(g) < 2:
            continue
        g.sort(key=lambda x: (x[1], x[2], x[3]))
        keep = g[0][0]
        for k, *_ , status in g[1:]:
            # never delete something already acted on
            if status in ("queued", "rejected"):
                continue
            drop.append(k)
    if drop and not dry:
        con.executemany("DELETE FROM jobs WHERE key=?", [(k,) for k in drop])
        con.commit()
    return len(drop)


# ----------------------------------------------------------------------- queue

def dedupe_queue(dry=False):
    items = {}
    for f in sorted(glob.glob(os.path.join(QUEUE, "*.json"))):
        if os.path.basename(f).startswith("_"):
            continue
        try:
            d = json.load(open(f))
        except json.JSONDecodeError:
            continue
        items.setdefault(d.get("company_slug") or d.get("id"), []).append((f, d))
    removed = 0
    for slug, lst in items.items():
        if len(lst) < 2:
            continue
        # keep the furthest-along item; on a tie keep the one with most fields
        lst.sort(key=lambda fd: (RANK.get(fd[1].get("status"), 0),
                                 len(fd[1].get("fields") or {}),
                                 fd[1].get("created", "")), reverse=True)
        for f, d in lst[1:]:
            if d.get("status") in SENT or d.get("submit_presses"):
                continue
            if not dry:
                for ext in (".json", ".png", ".html"):
                    p = f[:-5] + ext
                    if os.path.isfile(p):
                        os.remove(p)
            removed += 1
    return removed


# -------------------------------------------------------------------- trackers

def _dedupe_xlsx(path, key_cols, prefer_col=None, dry=False):
    if not os.path.isfile(path):
        return 0
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(hdr) if h}
    if not all(c in idx for c in key_cols):
        return 0
    seen, drop = {}, []
    for r in range(2, ws.max_row + 1):
        k = tuple(norm(str(ws.cell(r, idx[c]).value or ""))[:60] for c in key_cols)
        if not any(k):
            continue
        if k in seen:
            # keep whichever row carries a real status/date, drop the emptier one
            prev = seen[k]
            a = str(ws.cell(prev, idx[prefer_col]).value or "") if prefer_col else ""
            b = str(ws.cell(r, idx[prefer_col]).value or "") if prefer_col else ""
            if len(b) > len(a):
                drop.append(prev)
                seen[k] = r
            else:
                drop.append(r)
        else:
            seen[k] = r
    if drop and not dry:
        for r in sorted(drop, reverse=True):
            ws.delete_rows(r)
        wb.save(path)
    return len(drop)


def dedupe_tracker(dry=False):
    return _dedupe_xlsx(os.path.join(TRACKERS, "job-tracker.xlsx"),
                        ["Company", "Role"], "Stage", dry)


def dedupe_referrals(dry=False):
    return _dedupe_xlsx(os.path.join(TRACKERS, "referral-tracker.xlsx"),
                        ["Person", "Role Applied"], "Status", dry)


STEPS = {"boards": dedupe_boards, "jobs": dedupe_jobs, "queue": dedupe_queue,
         "tracker": dedupe_tracker, "referrals": dedupe_referrals}


def run_all(dry=False, only=None):
    total = 0
    for name, fn in STEPS.items():
        if only and name != only:
            continue
        try:
            n = fn(dry)
        except Exception as e:
            print(f"  {name:<10} FAILED {type(e).__name__}: {e}")
            continue
        total += n
        print(f"  {name:<10} {'would remove' if dry else 'removed'} {n}")
    return total


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--only", choices=sorted(STEPS))
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(line_buffering=True)
    except AttributeError:
        pass
    n = run_all(a.dry_run, a.only)
    print(f"  {'would remove' if a.dry_run else 'removed'} {n} duplicate(s) in total")


if __name__ == "__main__":
    main()
