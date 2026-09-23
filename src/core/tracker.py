#!/usr/bin/env python3
"""
tracker.py — keep tracking/job-tracker.xlsx in sync with the pipeline (Phase 6).

Writes into the tracker Sunil already uses, matching its existing columns and
Stage vocabulary, so his manual rows and the pipeline's rows live together.

  - queued/ranked jobs      -> Stage "To Apply"
  - submitted applications  -> Stage "Applied", with Date Applied + Follow-up Date
  - follow-ups due          -> listed, and surfaced in the daily digest

Follow-up is 5 BUSINESS days after applying, per job-tracker-guide.md.

Usage:
    python jobs/tracker.py --sync          # add new rows, update stages
    python jobs/tracker.py --followups     # what is due today
    python jobs/tracker.py --dry-run
"""
import argparse
import datetime as dt
import glob
import json
import os
import sqlite3
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from jobpilot.core.paths import (SRC as JOBS_DIR, TOOL, CONFIG, DATA, TRACKING, POLICY,  # noqa: E402
                   RESUME, APPLICATIONS, TRACKERS, MEMORY)
from jobpilot.core.config import cfg  # noqa: E402
XLSX = os.path.join(TRACKERS, "job-tracker.xlsx")
QUEUE_DIR = os.path.join(DATA, "queue")
DB = os.path.join(DATA, "state.db")

REGION = {"uae": "UAE", "saudi": "Saudi", "netherlands": "Netherlands",
          "australia": "Australia", "usa": "US", "remote_india": "Remote",
          "default": "Other"}


def business_days_from(d, n=None):
    n = n if n is not None else cfg("followups.job_days", 5)
    cur, added = d, 0
    while added < n:
        cur += dt.timedelta(days=1)
        if cur.weekday() < 5:
            added += 1
    return cur


def load_sheet():
    wb = load_workbook(XLSX)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(hdr) if h}
    return wb, ws, idx


def existing_keys(ws, idx):
    keys = set()
    for r in range(2, ws.max_row + 1):
        co = ws.cell(r, idx["Company"]).value
        role = ws.cell(r, idx["Role"]).value
        if co:
            keys.add((str(co).strip().lower(), str(role or "").strip().lower()[:40]))
    return keys


def row_for(ws, idx, co, role):
    for r in range(2, ws.max_row + 1):
        if (str(ws.cell(r, idx["Company"]).value or "").strip().lower() == co.lower()
                and str(ws.cell(r, idx["Role"]).value or "").strip().lower()[:40]
                == role.lower()[:40]):
            return r
    return None


def submitted_items():
    out = []
    for p in sorted(glob.glob(os.path.join(QUEUE_DIR, "*.json"))):
        if os.path.basename(p).startswith("_"):
            continue
        try:
            d = json.load(open(p))
        except json.JSONDecodeError:
            continue
        if d.get("status") in ("submitted", "approved", "skipped", "failed"):
            out.append(d)
    return out


def sync(dry=False, top=8):
    if not os.path.isfile(XLSX):
        sys.exit(f"tracker not found: {XLSX}")
    wb, ws, idx = load_sheet()
    have = existing_keys(ws, idx)
    added = updated = 0
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def put(r, col, val):
        if col in idx and val is not None:
            c = ws.cell(r, idx[col], val)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border

    # 1) applications the pipeline actually acted on
    for d in submitted_items():
        co, role = (d.get("company") or "?").strip(), (d.get("role") or "?").strip()
        stage = {"submitted": "Applied", "approved": "To Apply",
                 "skipped": "Withdrawn", "failed": "To Apply"}[d["status"]]
        r = row_for(ws, idx, co, role)
        if r is None:
            r = ws.max_row + 1
            added += 1
        else:
            updated += 1
        # Register it NOW, or the ranked-jobs pass below re-adds the same role.
        have.add((co.lower(), role.lower()[:40]))
        put(r, "Company", co)
        put(r, "Role", role)
        put(r, "Location", d.get("location"))
        put(r, "Region", REGION.get(d.get("market"), "Other"))
        put(r, "Work Mode", "Remote" if "remote" in (d.get("location") or "").lower()
            else "On-site")
        put(r, "Source", d.get("portal"))
        put(r, "Job Link", d.get("url"))
        put(r, "Expected/Posted Salary", d.get("salary_quoted"))
        put(r, "Resume Version", os.path.basename(d.get("resume") or ""))
        put(r, "Stage", stage)
        if d["status"] == "submitted" and d.get("submitted_at"):
            day = dt.date.fromisoformat(d["submitted_at"][:10])
            put(r, "Date Applied", day.isoformat())
            put(r, "Follow-up Date", business_days_from(day, 5).isoformat())
            put(r, "Next Action", "Follow up if no reply")
        elif d["status"] == "skipped":
            put(r, "Next Action", d.get("skip_reason", "skipped"))
        else:
            put(r, "Next Action", "Approve + submit")
        put(r, "Priority", "High" if (d.get("score") or 0) >= 60 else "Medium")

    # 2) the best ranked jobs not yet acted on
    if os.path.isfile(DB):
        con = sqlite3.connect(DB)
        rows = con.execute(
            "SELECT company,title,location,url,market,score,source FROM jobs "
            "WHERE market != 'usa' AND score >= 50 ORDER BY score DESC LIMIT ?",
            (top,)).fetchall()
        for co, title, loc, url, mk, score, src in rows:
            key = (co.strip().lower(), title.strip().lower()[:40])
            if key in have:
                continue
            have.add(key)
            r = ws.max_row + 1
            added += 1
            put(r, "Company", co); put(r, "Role", title); put(r, "Location", loc)
            put(r, "Region", REGION.get(mk, "Other"))
            put(r, "Work Mode", "Remote" if "remote" in (loc or "").lower() else "On-site")
            put(r, "Source", src); put(r, "Job Link", url)
            put(r, "Stage", "To Apply")
            put(r, "Next Action", "Tailor + fill")
            put(r, "Priority", "High" if score >= 60 else "Medium")
            put(r, "Notes", f"pipeline score {score}")

    if dry:
        print(f"  [dry-run] would add {added}, update {updated}")
        return added, updated
    wb.save(XLSX)
    print(f"  tracker: +{added} new, {updated} updated -> {XLSX}")
    return added, updated


def followups(days_ahead=0):
    if not os.path.isfile(XLSX):
        return []
    wb, ws, idx = load_sheet()
    today = dt.date.today()
    due = []
    for r in range(2, ws.max_row + 1):
        raw = ws.cell(r, idx.get("Follow-up Date", 12)).value
        stage = str(ws.cell(r, idx["Stage"]).value or "")
        if not raw or stage in ("Rejected", "Withdrawn", "Accepted", "Ghosted"):
            continue
        try:
            d = raw.date() if hasattr(raw, "date") else dt.date.fromisoformat(str(raw)[:10])
        except Exception:
            continue
        if d <= today + dt.timedelta(days=days_ahead):
            due.append({"company": ws.cell(r, idx["Company"]).value,
                        "role": ws.cell(r, idx["Role"]).value,
                        "due": d.isoformat(), "stage": stage,
                        "overdue": (today - d).days})
    return sorted(due, key=lambda x: x["due"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sync", action="store_true")
    ap.add_argument("--followups", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(line_buffering=True)
    except AttributeError:
        pass

    if a.followups:
        due = followups()
        if not due:
            print("  nothing due")
        for d in due:
            tag = f"{d['overdue']}d overdue" if d["overdue"] > 0 else "due today"
            print(f"  [{tag:>12}] {d['company']} — {d['role']}  ({d['stage']})")
        return
    sync(dry=a.dry_run)


if __name__ == "__main__":
    main()
