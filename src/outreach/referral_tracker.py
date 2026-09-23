#!/usr/bin/env python3
"""
referral_tracker.py — turn a submitted application into referral outreach.

An application sitting in an ATS queue is the weakest version of applying. This
picks 3-4 people worth contacting at that company, drafts a message for each,
records them in a tracker Sunil can work through, and pushes them to Telegram in
copy-paste form.

Candidates are a deliberate mix:
  * 1st-degree connections (from the LinkedIn export) — warm, ask directly
  * cold but well-matched people (GitHub org, paper authors) — invite first

Sending is ALWAYS manual. This produces text and links; Sunil presses send.

Sheet: tracking/referral-tracker.xlsx
Status flow: To Contact -> Invite Sent -> Accepted -> Message Sent -> Replied
             -> Referred | No Response

Usage:
    python jobs/referral_tracker.py --for <company-slug>   # from an application
    python jobs/referral_tracker.py --company Databricks --role "AI Engineer"
    python jobs/referral_tracker.py --list                 # everything open
    python jobs/referral_tracker.py --due                  # follow-ups due today
    python jobs/referral_tracker.py --set <row> <status>
    python jobs/referral_tracker.py --digest               # evening summary
"""
import argparse
import datetime as dt
import glob
import json
import os
import sys
import urllib.parse
import urllib.request

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from jobpilot.core.paths import (SRC as JOBS_DIR, TOOL, CONFIG, DATA, TRACKING, POLICY,  # noqa: E402
                   RESUME, APPLICATIONS, TRACKERS, MEMORY)
from jobpilot.core.config import cfg  # noqa: E402
XLSX = os.path.join(TRACKERS, "referral-tracker.xlsx")
ENV = os.path.expanduser("~/.config/jobbot/env")

COLS = ["Person", "Their Title", "Company", "Role Applied", "Relationship",
        "Why Them", "Profile URL", "Status", "Invite Sent", "Message Sent",
        "Follow-up Date", "Replied", "Message Used", "Notes"]

STATUSES = ["To Contact", "Invite Sent", "Accepted", "Message Sent", "Replied",
            "Referred", "No Response"]

# Chase an unanswered invite once, after a week. Sooner reads as pushy; later and
# they have forgotten the context entirely.
FOLLOWUP_DAYS = cfg("followups.referral_days", 7)


# --------------------------------------------------------------------- sheet

def ensure_sheet():
    if os.path.isfile(XLSX):
        return load_workbook(XLSX)
    wb = Workbook()
    ws = wb.active
    ws.title = "Referrals"
    ws.append(COLS)
    head = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="D9D9D9")
    for c in ws[1]:
        c.fill = head
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, w in zip("ABCDEFGHIJKLMN",
                      (22, 30, 18, 28, 14, 24, 38, 14, 12, 12, 13, 10, 60, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    wb.save(XLSX)
    return wb


def rows():
    wb = ensure_sheet()
    ws = wb.active
    out = []
    for r in range(2, ws.max_row + 1):
        d = {c: ws.cell(r, i + 1).value for i, c in enumerate(COLS)}
        if d["Person"]:
            d["_row"] = r
            out.append(d)
    return out


def add(entries):
    """Append candidates, skipping anyone already tracked for the same role."""
    wb = ensure_sheet()
    ws = wb.active
    have = {(str(ws.cell(r, 1).value or "").lower(),
             str(ws.cell(r, 4).value or "").lower())
            for r in range(2, ws.max_row + 1)}
    thin = Side(style="thin", color="D9D9D9")
    added = 0
    for e in entries:
        key = (e["Person"].lower(), e["Role Applied"].lower())
        if key in have:
            continue
        have.add(key)
        ws.append([e.get(c, "") for c in COLS])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        added += 1
    wb.save(XLSX)
    return added


def set_status(row, status):
    if status not in STATUSES:
        sys.exit(f"status must be one of: {', '.join(STATUSES)}")
    wb = ensure_sheet()
    ws = wb.active
    today = dt.date.today()
    ws.cell(row, COLS.index("Status") + 1, status)
    if status == "Invite Sent":
        ws.cell(row, COLS.index("Invite Sent") + 1, today.isoformat())
        ws.cell(row, COLS.index("Follow-up Date") + 1,
                (today + dt.timedelta(days=FOLLOWUP_DAYS)).isoformat())
    elif status == "Message Sent":
        ws.cell(row, COLS.index("Message Sent") + 1, today.isoformat())
        ws.cell(row, COLS.index("Follow-up Date") + 1,
                (today + dt.timedelta(days=FOLLOWUP_DAYS)).isoformat())
    elif status in ("Replied", "Referred", "No Response"):
        ws.cell(row, COLS.index("Replied") + 1,
                today.isoformat() if status == "Replied" else "")
        ws.cell(row, COLS.index("Follow-up Date") + 1, "")
    wb.save(XLSX)
    print(f"  row {row} -> {status}")


def due(days_ahead=0):
    today = dt.date.today()
    out = []
    for d in rows():
        if d["Status"] in ("Referred", "No Response", "Replied"):
            continue
        raw = d["Follow-up Date"]
        if not raw:
            continue
        try:
            when = raw.date() if hasattr(raw, "date") else dt.date.fromisoformat(str(raw)[:10])
        except Exception:
            continue
        if when <= today + dt.timedelta(days=days_ahead):
            d["_overdue"] = (today - when).days
            out.append(d)
    return sorted(out, key=lambda x: -x["_overdue"])


# ---------------------------------------------------------------- candidates

def candidates_for(company, role, want=4):
    """A mix of warm and cold, warm first.

    Always leaves room for at least one cold candidate: a warm list alone gives
    nowhere to go when nobody replies, and cold targets are where new network
    gets built.
    """
    from jobpilot.outreach import outreach
    out = []
    warm_cap = max(1, want - 1)

    try:
        from jobpilot.outreach import referrals
        conns = referrals.load_connections()
        if conns:
            for pts, why, c in referrals.match(conns, [company]).get(company, [])[:warm_cap]:
                name = f"{c['first']} {c['last']}".strip()
                out.append({
                    "Person": name, "Their Title": c["position"][:60],
                    "Company": company, "Role Applied": role,
                    "Relationship": "1st-degree", "Why Them": why,
                    "Profile URL": c["url"], "Status": "To Contact",
                    "Message Used": outreach.referral_ask(
                        company, role, c["first"], c["position"][:40]),
                })
    except Exception as e:
        print(f"  [warn] connections lookup: {type(e).__name__}")

    if len(out) < want:
        try:
            from jobpilot.outreach import prospects
            gh = company.lower().replace(" ", "")
            people, _ = prospects.github_people(gh, limit=6)
            if not people:
                people, _ = prospects.github_commit_authors(gh, limit=6)
            for p in people:
                if len(out) >= want:
                    break
                out.append({
                    "Person": p["name"][:40], "Their Title": (p.get("bio") or "Engineer")[:60],
                    "Company": company, "Role Applied": role,
                    "Relationship": "cold (GitHub)",
                    "Why Them": "engineer on public repos",
                    "Profile URL": p.get("url", ""), "Status": "To Contact",
                    "Message Used": outreach.connection_note(company, role),
                })
        except Exception as e:
            print(f"  [warn] github lookup: {type(e).__name__}")

    if len(out) < want:
        try:
            from jobpilot.outreach import prospects
            papers, _ = prospects.openalex_people(company, limit=6)
            for p in papers:
                if len(out) >= want:
                    break
                title = p["papers"][0] if p["papers"] else ""
                out.append({
                    "Person": p["name"][:40], "Their Title": "Researcher",
                    "Company": company, "Role Applied": role,
                    "Relationship": "cold (paper)",
                    "Why Them": f"published: {title[:50]}",
                    "Profile URL": p.get("oa", ""), "Status": "To Contact",
                    "Message Used": outreach.paper_author(
                        company, role, p["name"].split()[0], title),
                })
        except Exception as e:
            print(f"  [warn] openalex lookup: {type(e).__name__}")

    return out[:want]


# ------------------------------------------------------------------ telegram

def env():
    d = {}
    if os.path.isfile(ENV):
        for line in open(ENV):
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                d[k.strip()] = v.strip()
    return d


def telegram(text):
    e = env()
    tok, chat = e.get("TELEGRAM_BOT_TOKEN"), e.get("TELEGRAM_CHAT_ID")
    if not (tok and chat):
        return False
    data = urllib.parse.urlencode({"chat_id": chat, "text": text[:4000],
                                   "parse_mode": "HTML",
                                   "disable_web_page_preview": "true"}).encode()
    try:
        with urllib.request.urlopen(
                f"https://api.telegram.org/bot{tok}/sendMessage", data, timeout=20) as r:
            return json.load(r).get("ok", False)
    except Exception:
        return False


def esc(s):
    return (str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def referral_link():
    """One page over Tailscale beats a dozen Telegram messages on a phone."""
    e = env()
    tok = e.get("FORM_TOKEN", "")
    host = ""
    try:
        import subprocess
        out = subprocess.run(["tailscale", "status", "--json"],
                             capture_output=True, timeout=6).stdout
        host = json.loads(out).get("Self", {}).get("DNSName", "").rstrip(".")
    except Exception:
        pass
    if not host:
        import socket
        s_ = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        try:
            s_.connect(("8.8.8.8", 80))
            host = s_.getsockname()[0]
        except Exception:
            host = "localhost"
        finally:
            s_.close()
    return f"http://{host}:{cfg('server.port', 8765)}/referrals" + (f"?t={tok}" if tok else "")


def push(company, role, entries):
    """ONE message with ONE link. Everything is on the page: profile, message to
    copy, and status buttons that write back to the tracker."""
    warm = sum(1 for e in entries if e["Relationship"] == "1st-degree")
    names = ", ".join(e["Person"].split()[0] for e in entries[:4])
    body = (f"🤝 <b>Referral targets — {esc(company)}</b>\n"
            f"<i>{esc(role)}</i>\n\n"
            f"{len(entries)} people ({warm} you already know): {esc(names)}\n\n"
            f"Open profiles, copy the drafted message, set status — all on one page:\n"
            f"{referral_link()}")
    return telegram(body)


# ---------------------------------------------------------------------- main

def from_application(slug):
    """Read company + role off the queue item for a submitted application."""
    fs = sorted(glob.glob(os.path.join(DATA, "queue", f"{slug}-*.json")))
    if not fs:
        sys.exit(f"no queue item for {slug}")
    d = json.load(open(fs[-1]))
    return d.get("company", slug), d.get("role", "the role")


def digest():
    open_rows = [d for d in rows() if d["Status"] not in ("Referred", "No Response")]
    d_due = due()
    lines = [f"🤝 <b>Referrals</b> — {dt.date.today():%a %d %b}", ""]
    if d_due:
        lines.append(f"<b>⏰ {len(d_due)} to chase</b>")
        for d in d_due[:6]:
            tag = f"{d['_overdue']}d late" if d["_overdue"] > 0 else "today"
            lines.append(f"• {esc(d['Person'])} — {esc(d['Company'])} ({d['Status']}, {tag})")
        lines.append("")
    todo = [d for d in open_rows if d["Status"] == "To Contact"]
    if todo:
        lines.append(f"<b>{len(todo)} not contacted yet</b>")
        for d in todo[:6]:
            star = "★ " if d["Relationship"] == "1st-degree" else ""
            lines.append(f"• {star}{esc(d['Person'])} — {esc(d['Company'])}")
    if not d_due and not todo:
        lines.append("Nothing outstanding.")
    lines += ["", f"Work through them here:\n{referral_link()}",
              f"<i>Sheet: tracking/referral-tracker.xlsx ({len(open_rows)} open)</i>"]
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--for", dest="slug", help="application slug to build referrals for")
    ap.add_argument("--company")
    ap.add_argument("--role", default="the role")
    ap.add_argument("--want", type=int, default=4)
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--due", action="store_true")
    ap.add_argument("--digest", action="store_true")
    ap.add_argument("--no-telegram", action="store_true")
    ap.add_argument("--set", nargs=2, metavar=("ROW", "STATUS"))
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(line_buffering=True)
    except AttributeError:
        pass

    if a.set:
        return set_status(int(a.set[0]), a.set[1])
    if a.list:
        for d in rows():
            star = "★" if d["Relationship"] == "1st-degree" else " "
            print(f"  {d['_row']:>3} {star} {str(d['Person'])[:22]:<24} "
                  f"{str(d['Company'])[:14]:<16} {str(d['Status']):<12} "
                  f"{str(d['Follow-up Date'] or '')[:10]}")
        return
    if a.due:
        for d in due():
            tag = f"{d['_overdue']}d late" if d["_overdue"] > 0 else "today"
            print(f"  [{tag:>8}] row {d['_row']}  {d['Person']} — {d['Company']} ({d['Status']})")
        return
    if a.digest:
        t = digest()
        print(t.replace("<b>", "").replace("</b>", "").replace("<i>", "").replace("</i>", ""))
        if not a.no_telegram:
            telegram(t)
        return

    company, role = (from_application(a.slug) if a.slug
                     else (a.company, a.role))
    if not company:
        ap.error("give --for <slug> or --company")
    print(f"  finding referral targets: {company} — {role}")
    ents = candidates_for(company, role, a.want)
    if not ents:
        print("  none found")
        return
    n = add(ents)
    print(f"  {len(ents)} candidate(s), {n} new in the tracker")
    for e in ents:
        print(f"    [{e['Relationship']:<14}] {e['Person'][:26]:<28} {e['Their Title'][:40]}")
    if not a.no_telegram:
        print("  telegram:", "sent" if push(company, role, ents) else "FAILED")


if __name__ == "__main__":
    main()
